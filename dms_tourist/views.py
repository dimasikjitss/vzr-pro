from datetime import date, timedelta
from decimal import Decimal
import json

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db import transaction

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill

from insurance.models import Branch, IndividualHolder, CompanyHolder
from .models import DMSTouristPolicy, DMSTouristInsuredPerson, DMSTouristTariff
from django.core.serializers.json import DjangoJSONEncoder
from .services import calculate_single_person_premium


def is_superadmin(user):
    return user.is_superuser


from django.db import transaction

@login_required
def create_policy(request):
    """Оформление и сохранение полиса ДМС Турист (индивидуальный или групповой) с привязкой страхователя к справочнику"""
    # 1. Проверка филиала агента
    try:
        user_branch = request.user.profile.branch
    except Exception:
        return render(request, 'insurance/error.html', {'message': 'Филиал не настроен в профиле пользователя.'})
    

    if request.method == "POST":
        try:
            with transaction.atomic():
                holder_kind = request.POST.get('holder_kind', 'individual')
                s_date_str = request.POST.get('start_date')
                e_date_str = request.POST.get('end_date')
                sum_insured = int(request.POST.get('sum_insured', 100000))
                purpose = request.POST.get('purpose', 'BUSINESS')

                d_from = date.fromisoformat(s_date_str)
                d_to = date.fromisoformat(e_date_str)
                days_count = (d_to - d_from).days + 1

                if days_count < 1 or days_count > 365:
                    raise ValueError("Период страхования должен составлять от 1 до 365 дней.")

                # Получение базового тарифа из БД
                tariff = DMSTouristTariff.objects.filter(
                    sum_insured=sum_insured,
                    days_min__lte=days_count,
                    days_max__gte=days_count
                ).first()

                base_rate = Decimal(str(tariff.base_daily_rate)) if tariff else Decimal('80.00')
                k_dur = Decimal(str(tariff.duration_coeff)) if tariff else Decimal('1.00')

                # Коэффициент цели поездки
                k_pur = Decimal('1.00')
                if purpose == 'TREKKING': 
                    k_pur = Decimal('1.50')
                elif purpose == 'SPORT': 
                    k_pur = Decimal('2.00')

                holder_instance = None
                
                # 2. Определение и создание страхователя в справочнике (как в Номаде)
                if holder_kind == 'company':
                    h_name = request.POST.get('company_name', '').strip().upper()
                    h_inn = request.POST.get('company_inn', '').strip()
                    h_phone = request.POST.get('company_phone', '').strip()
                    h_addr = request.POST.get('company_address', '').strip().upper()
                    c_okpo = request.POST.get('company_okpo', '').strip()
                    c_sign = request.POST.get('company_signatory', '').strip().upper()
                    
                    holder_instance = CompanyHolder.objects.create(
                        company_name=h_name,
                        inn=h_inn,
                        okpo=c_okpo,
                        legal_address=h_addr,
                        signatory_name=c_sign,
                        phone=h_phone,
                    )
                    h_pass = h_inn
                    h_citizenship = ''
                    h_relative = c_sign
                else:
                    h_name = request.POST.get('holder_name', '').strip().upper()
                    h_citizenship = request.POST.get('holder_citizenship', '').strip()
                    h_pass = request.POST.get('holder_passport', '').strip().upper()
                    h_inn = (request.POST.get('holder_inn', '') or request.POST.get('holder_pin', '')).strip()
                    h_phone = request.POST.get('holder_phone', '').strip()
                    h_addr = request.POST.get('holder_address', '').strip().upper()
                    h_relative = request.POST.get('holder_relative', '').strip() or "Не указан"
                    c_okpo = ''
                    c_sign = ''

                    holder_instance = IndividualHolder.objects.create(
                        name=h_name,
                        inn=h_inn,
                        phone=h_phone,
                        address=h_addr,
                        passport=h_pass,
                        relative_contact=h_relative,
                    )

                # 3. Сбор параметров для создания полиса с передачей объекта holder
                policy_kwargs = {
                    'agent': request.user,
                    'branch': user_branch,
                    'holder': holder_instance,  # Заполняет GenericForeignKey (holder_type и holder_id)
                    'holder_name': h_name,
                    'holder_citizenship': h_citizenship,
                    'holder_passport': h_pass,
                    'holder_inn': h_inn,
                    'holder_phone': h_phone,
                    'holder_address': h_addr,
                    'relative_contact': h_relative,
                    'date_from': d_from,
                    'date_to': d_to,
                    'days_count': days_count,
                    'sum_insured': sum_insured,
                    'purpose': purpose,
                }

                if hasattr(DMSTouristPolicy, 'company_okpo'):
                    policy_kwargs['company_okpo'] = c_okpo
                if hasattr(DMSTouristPolicy, 'company_signatory'):
                    policy_kwargs['company_signatory'] = c_sign

                policy = DMSTouristPolicy.objects.create(**policy_kwargs)

                # 4. Обработка списка застрахованных туристов
                names = request.POST.getlist('insured_name[]') or request.POST.getlist('insured_name')
                dobs = request.POST.getlist('insured_dob[]') or request.POST.getlist('insured_dob')
                citizenships = request.POST.getlist('insured_citizenship[]') or request.POST.getlist('insured_citizenship')
                passports = request.POST.getlist('insured_passport[]') or request.POST.getlist('insured_passport')

                total_netto = Decimal('0.00')
                persons_to_create = []

                for i in range(len(names)):
                    n_val = names[i].strip().upper()
                    if not n_val:
                        continue
                    
                    c_val = citizenships[i].strip() if i < len(citizenships) else ''
                    p_val = passports[i].strip().upper() if i < len(passports) else ''
                    dob_raw = dobs[i].strip() if i < len(dobs) else ''

                    if not dob_raw:
                        raise ValueError(f"Не указана дата рождения для застрахованного {n_val}")

                    if '-' in dob_raw:
                        dob_date = date.fromisoformat(dob_raw)
                    elif '.' in dob_raw:
                        parts = dob_raw.split('.')
                        dob_date = date(int(parts[2]), int(parts[1]), int(parts[0]))
                    else:
                        dob_date = date(1990, 1, 1)

                    calc_res = calculate_single_person_premium(
                        sum_insured=sum_insured,
                        days=days_count,
                        birth_date=dob_date,
                        start_date=d_from,
                        purpose=purpose
                    )

                    total_netto += calc_res['netto']

                    persons_to_create.append(DMSTouristInsuredPerson(
                        policy=policy,
                        full_name=n_val,
                        birth_date=dob_date,
                        citizenship=c_val,
                        passport_number=p_val,
                        age=calc_res['age'],
                        age_coeff=calc_res['k_age'],
                        duration_coeff=calc_res['k_dur'],
                        purpose_coeff=calc_res['k_pur'],
                        premium_netto=calc_res['netto'],
                        nsp=calc_res['nsp'],
                        premium_gross=calc_res['gross']
                    ))

                if not persons_to_create:
                    raise ValueError("Укажите хотя бы одного застрахованного туриста.")

                DMSTouristInsuredPerson.objects.bulk_create(persons_to_create)

                # 5. Пересчет и сохранение итогов полиса
                total_nsp = (total_netto * Decimal('0.02')).quantize(Decimal('0.01'))
                total_gross = total_netto + total_nsp

                policy.total_premium_netto = total_netto
                policy.total_nsp = total_nsp
                policy.total_premium_gross = total_gross
                policy.save(update_fields=['total_premium_netto', 'total_nsp', 'total_premium_gross'])

                messages.success(request, f"Полис {policy.policy_number} успешно оформлен!")
                return redirect('dms_tourist:policy_success', pk=policy.pk)

        except Exception as e:
            messages.error(request, f"Ошибка при сохранении полиса: {e}")
            return HttpResponse(f"Ошибка сохранения полиса: {str(e)}", status=400)

    # GET-запрос
    tariffs = DMSTouristTariff.objects.all().order_by('sum_insured', 'days_min')
    return render(request, 'dms_tourist/create.html', {
        'tariffs': tariffs,
        'today': date.today().isoformat()
    })

@login_required
def policy_preview(request):
    """Предпросмотр черновика полиса ДМС Турист (до сохранения в БД)"""
    if request.method != "POST":
        return redirect('dms_tourist:create_policy')

    try:
        holder_kind = request.POST.get('holder_kind', 'individual')
        s_date_str = request.POST.get('start_date')
        e_date_str = request.POST.get('end_date')
        sum_insured = int(request.POST.get('sum_insured', 100000))
        purpose = request.POST.get('purpose', 'BUSINESS')

        d_from = date.fromisoformat(s_date_str)
        d_to = date.fromisoformat(e_date_str)
        days_count = (d_to - d_from).days + 1

        if days_count < 1 or days_count > 365:
            raise ValueError("Период страхования должен составлять от 1 до 365 дней.")

        # Поиск базового тарифа в БД
        tariff = DMSTouristTariff.objects.filter(
            sum_insured=sum_insured,
            days_min__lte=days_count,
            days_max__gte=days_count
        ).first()

        base_rate = Decimal(str(tariff.base_daily_rate)) if tariff else Decimal('80.00')
        k_dur = Decimal(str(tariff.duration_coeff)) if tariff else Decimal('1.00')

        # Коэффициент цели поездки
        k_pur = Decimal('1.00')
        # if purpose == 'TREKKING': 
        #     k_pur = Decimal('1.50')
        # elif purpose == 'SPORT': 
        #     k_pur = Decimal('2.00')
        # Для BUSINESS, STUDENT, RESIDENCE, REGISTRATION коэффициент цели равен 1.00

        # Списки застрахованных туристов из формы (включая гражданство)
        names = request.POST.getlist('insured_name[]') or request.POST.getlist('insured_name')
        dobs = request.POST.getlist('insured_dob[]') or request.POST.getlist('insured_dob')
        h_citizenship = request.POST.get('holder_citizenship', '').strip()
        citizenships = request.POST.getlist('insured_citizenship[]') or request.POST.getlist('insured_citizenship')
        passports = request.POST.getlist('insured_passport[]') or request.POST.getlist('insured_passport')

        insured_persons_list = []
        total_netto = Decimal('0.00')
        has_elderly = False

        for i in range(len(names)):
            n_val = names[i].strip().upper()
            if not n_val:
                continue
            
            c_val = citizenships[i].strip() if i < len(citizenships) else ''
            p_val = passports[i].strip().upper() if i < len(passports) else ''
            dob_raw = dobs[i].strip() if i < len(dobs) else ''

            dob_date = None
            age = 30
            if dob_raw:
                try:
                    if '-' in dob_raw:
                        dob_date = date.fromisoformat(dob_raw)
                    elif '.' in dob_raw:
                        parts = dob_raw.split('.')
                        dob_date = date(int(parts[2]), int(parts[1]), int(parts[0]))
                    age = d_from.year - dob_date.year - ((d_from.month, d_from.day) < (dob_date.month, dob_date.day))
                except Exception:
                    dob_date = date(1990, 1, 1)
                    age = 30

            # Проверка возрастных лимитов (от 1 до 75 лет)
            if age < 1 or age > 75:
                raise ValueError(f"Возраст застрахованного {n_val} ({age} лет) выходит за допустимые пределы (от 1 до 75 лет).")

            if age >= 61:
                has_elderly = True

            k_age = Decimal('1.00')
            p_netto = (Decimal(days_count) * base_rate * k_dur * k_age * k_pur).quantize(Decimal('0.01'))
            p_nsp = (p_netto * Decimal('0.02')).quantize(Decimal('0.01'))
            p_gross = p_netto + p_nsp

            total_netto += p_netto

            insured_persons_list.append({
                'full_name': n_val,
                'birth_date': dob_date,
                'age': age,
                
                'citizenship': c_val,  # Передаем гражданство в шаблон предпросмотра
                'passport_number': p_val,
                'premium_netto': p_netto,
                'nsp': p_nsp,
                'premium_gross': p_gross,
            })

        total_nsp = (total_netto * Decimal('0.02')).quantize(Decimal('0.01'))
        total_gross = total_netto + total_nsp

        preview_policy = {
            'policy_number': 'ПРОЕКТ / DRAFT',
            'issue_date': date.today(),
            'agent': request.user,
            'holder_kind': holder_kind,
            'holder_citizenship': h_citizenship,
            'holder_name': request.POST.get('company_name', '').strip().upper() if holder_kind == 'company' else request.POST.get('holder_name', '').strip().upper(),
            'holder_passport': request.POST.get('company_inn', '').strip() if holder_kind == 'company' else request.POST.get('holder_passport', '').strip().upper(),
            'holder_inn': request.POST.get('company_inn', '').strip() if holder_kind == 'company' else request.POST.get('holder_inn', '').strip(),
            'company_okpo': request.POST.get('company_okpo', '').strip(),
            'company_signatory': request.POST.get('company_signatory', '').strip().upper(),
            'holder_phone': request.POST.get('company_phone', '').strip() if holder_kind == 'company' else request.POST.get('holder_phone', '').strip(),
            'holder_address': request.POST.get('company_address', '').strip().upper() if holder_kind == 'company' else request.POST.get('holder_address', '').strip().upper(),
            'relative_contact': request.POST.get('holder_relative', '').strip() or '—',
            'date_from': d_from,
            'date_to': d_to,
            'days_count': days_count,
            'sum_insured': sum_insured,
            'purpose': purpose,
            'total_premium_netto': total_netto,
            'total_nsp': total_nsp,
            'total_premium_gross': total_gross,
        }

        return render(request, 'dms_tourist/print_policy.html', {
            'policy': preview_policy,
            'persons': insured_persons_list,
            'has_elderly': has_elderly,
            'is_draft': True
        })

    except Exception as e:
        return HttpResponse(f"Ошибка формирования предпросмотра: {str(e)}", status=400)
   
    
@login_required
def policy_success(request, pk):
    return render(request, 'dms_tourist/success.html', {'policy': get_object_or_404(DMSTouristPolicy, pk=pk)})


@login_required
def print_policy(request, pk):
    """Печать чистовика полиса ДМС Турист (1-й лист, 2-й лист правила, 3-й лист Приложение №1 при группе)"""
    policy = get_object_or_404(DMSTouristPolicy, pk=pk)
    persons = list(policy.insured_persons.all().order_by('id'))
    has_elderly = any(p.age >= 61 for p in persons)

    return render(request, 'dms_tourist/print_policy.html', {
        'policy': policy,
        'persons': persons,
        'has_elderly': has_elderly,
        'is_draft': False
    })

# --- 1. ЭКСПОРТ В БУХГАЛТЕРИЮ ---
@login_required
def export_excel_buh_dms(request):
    d_from_raw = request.GET.get('date_from')
    d_to_raw = request.GET.get('date_to')
    branch_id = request.GET.get('branch')

    try:
        d_from = date.fromisoformat(d_from_raw)
        d_to = date.fromisoformat(d_to_raw)
    except (TypeError, ValueError):
        d_from = d_to = date.today()

    filters = {'issue_date__range': [d_from, d_to]}
    if not request.user.is_superuser:
        profile = getattr(request.user, 'profile', None)
        if profile and profile.branch:
            filters['branch_id'] = profile.branch.id
        else:
            return HttpResponse("Вы не привязаны к филиалу", status=403)
    else:
        if branch_id and str(branch_id).strip():
            filters['branch_id'] = branch_id

    policies = DMSTouristPolicy.objects.filter(**filters).select_related('agent', 'branch').prefetch_related('insured_persons').order_by('issue_date', 'id')

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет ДМС Бухгалтерия"

    font_title = Font(name='Arial', size=14, bold=True, color="1E293B")
    font_bold = Font(name='Arial', size=10, bold=True)
    font_regular = Font(name='Arial', size=10, bold=False)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    fill_summary = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_group = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid") # Нежно-голубой для групп

    ws.append([f"Отчет по полисам ДМС Турист за период с {d_from.strftime('%d.%m.%Y')} по {d_to.strftime('%d.%m.%Y')}"])
    ws.row_dimensions[1].height = 35
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.cell(row=1, column=1).font = font_title
    ws.cell(row=1, column=1).alignment = align_center

    ws.append([])
    headers = [
        'Серия', 'Номер полиса', 'Дата оформления', 'Страхователь',
        'Застрахованный турист', 'Номер паспорта', 'Дата рождения', 'Телефон',
        'Дата въезда', 'Дата выезда', 'Страховая сумма, сом',
        'Премия без НСП, сом', 'НСП 2%', 'Стр. премия с НСП, сом', 'Кто оформил'
    ]
    ws.append(headers)
    ws.row_dimensions[3].height = 28
    for cell in ws[3]:
        cell.font = font_bold
        cell.alignment = align_center

    total_count = 0
    total_gross = Decimal('0.00')
    branch_stats = {}

    for p in policies:
        b_code = p.branch.code
        if b_code not in branch_stats:
            branch_stats[b_code] = {'name': p.branch.name, 'count': 0, 'sum': Decimal('0.00')}
        
        total_count += 1
        branch_stats[b_code]['count'] += 1

        persons = list(p.insured_persons.all())
        count_persons = len(persons)

        if count_persons == 0:
            continue
        elif count_persons == 1:
            # 1 человек — обычная строка
            person = persons[0]
            total_gross += person.premium_gross
            branch_stats[b_code]['sum'] += person.premium_gross
            
            ws.append([
                'DMST', p.policy_number, p.issue_date.strftime('%d.%m.%Y'), p.holder_name,
                person.full_name, person.passport_number, person.birth_date.strftime('%d.%m.%Y'),
                p.holder_phone, p.date_from.strftime('%d.%m.%Y'), p.date_to.strftime('%d.%m.%Y'),
                p.sum_insured, person.premium_netto, person.nsp, person.premium_gross,
                p.agent.get_full_name() or p.agent.username
            ])
            cur_r = ws.max_row
            ws.row_dimensions[cur_r].height = 20
            for c in ws[cur_r]:
                c.font = font_regular
                c.alignment = align_center
        else:
            # Группа людей: выводим итоговую строку полиса, а под ней скрытые детали с группировкой (+)
            netto_total = sum(pr.premium_netto for pr in persons)
            nsp_total = sum(pr.nsp for pr in persons)
            gross_total = sum(pr.premium_gross for pr in persons)
            
            total_gross += gross_total
            branch_stats[b_code]['sum'] += gross_total

            # Главная строка полиса (сводная по группе)
            ws.append([
                'DMST', p.policy_number, p.issue_date.strftime('%d.%m.%Y'), p.holder_name,
                f"📎 Группа (всего застраховано: {count_persons} чел.)", "—", "—",
                p.holder_phone, p.date_from.strftime('%d.%m.%Y'), p.date_to.strftime('%d.%m.%Y'),
                p.sum_insured, netto_total, nsp_total, gross_total,
                p.agent.get_full_name() or p.agent.username
            ])
            summary_r = ws.max_row
            ws.row_dimensions[summary_r].height = 22
            for c in ws[summary_r]:
                c.font = font_bold
                c.fill = fill_group
                c.alignment = align_center

            # Строки с деталями каждого туриста (подчиненные, сворачиваются под «+»)
            for person in persons:
                ws.append([
                    '', '', '', '',
                    f"   ↳ {person.full_name}", person.passport_number, person.birth_date.strftime('%d.%m.%Y'),
                    '', '', '', '',
                    person.premium_netto, person.nsp, person.premium_gross, ''
                ])
                sub_r = ws.max_row
                ws.row_dimensions[sub_r].outline_level = 1  # Уровень группировки
                ws.row_dimensions[sub_r].hidden = True      # Свернуто по умолчанию (нажимается «+»)
                ws.row_dimensions[sub_r].height = 18
                for c in ws[sub_r]:
                    c.font = font_regular
                    c.alignment = align_center

    ws.append([])
    ws.append(['ОБЩИЕ ИТОГИ ЗА ПЕРИОД'])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=3)
    ws.cell(row=ws.max_row, column=1).font = font_bold
    ws.cell(row=ws.max_row, column=1).fill = fill_summary

    ws.append(['Всего оформлено полисов:', f"{total_count} шт."])
    ws.append(['Общая сумма премий с НСП:', f"{total_gross} сом"])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

    filename = f"Report_DMS_Buhgalteriya_{d_from}_{d_to}.xlsx"
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

# --- 2. ЭКСПОРТ ДЛЯ БИШКЕК АССИСТ ---
@login_required
@user_passes_test(is_superadmin)
def export_excel_bishkek_assist(request):
    d_from_raw = request.GET.get('date_from')
    d_to_raw = request.GET.get('date_to')

    try:
        d_from = date.fromisoformat(d_from_raw)
        d_to = date.fromisoformat(d_to_raw)
    except (TypeError, ValueError):
        d_from = d_to = date.today()

    policies = DMSTouristPolicy.objects.filter(issue_date__range=[d_from, d_to]).select_related('agent', 'branch').prefetch_related('insured_persons').order_by('issue_date', 'id')

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет Бишкек Ассист"

    font_title = Font(name='Arial', size=14, bold=True, color="1E293B")
    font_bold = Font(name='Arial', size=10, bold=True)
    font_regular = Font(name='Arial', size=10, bold=False)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.append([f"Реестр полисов «ДМС Турист» для «Бишкек Ассист» за период с {d_from.strftime('%d.%m.%Y')} по {d_to.strftime('%d.%m.%Y')}"])
    ws.row_dimensions[1].height = 35
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    ws.cell(row=1, column=1).font = font_title
    ws.cell(row=1, column=1).alignment = align_center

    ws.append([])
    headers = [
        'Номер полиса', 'Дата оформления', 'Начало действия', 'Окончание действия',
        'Цель поездки / Риск', 'Страхователь', 'Застрахованный турист',
        'Номер паспорта', 'Дата рождения', 'Возраст', 'Телефон в КР', 'Адрес в КР / Отель',
        'Страховая сумма, сом', 'Премия без НСП, сом', 'Премия с НСП, сом', 'Кто оформил'
    ]
    ws.append(headers)
    ws.row_dimensions[3].height = 28
    for cell in ws[3]:
        cell.font = font_bold
        cell.alignment = align_center

    for p in policies:
        for person in p.insured_persons.all():
            ws.append([
                p.policy_number,
                p.issue_date.strftime('%d.%m.%Y'),
                p.date_from.strftime('%d.%m.%Y'),
                p.date_to.strftime('%d.%m.%Y'),
                p.get_purpose_display(),
                p.holder_name,
                person.full_name,
                person.passport_number,
                person.birth_date.strftime('%d.%m.%Y'),
                person.age,
                p.holder_phone,
                p.holder_address or "—",
                p.sum_insured,
                person.premium_netto,
                person.premium_gross,
                p.agent.get_full_name() or p.agent.username
            ])
            cur_r = ws.max_row
            ws.row_dimensions[cur_r].height = 20
            for c in ws[cur_r]:
                c.font = font_regular
                c.alignment = align_center

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

    filename = f"GSO_Bishkek_Assist_Report_{d_from}_{d_to}.xlsx"
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

def get_dms_tariffs_api(request):
    """API отдачи тарифов для JS (как в ВЗР get_tariffs_api)"""
    tariffs = list(DMSTouristTariff.objects.all().values(
        'sum_insured', 'days_min', 'days_max', 'duration_coeff', 'base_daily_rate'
    ))
    return JsonResponse(tariffs, safe=False)